"""XLSX export using the deployed Python runtime's existing Excel library."""
from datetime import datetime, timezone
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def solar_archive_xlsx(captures):
    book=Workbook();sheet=book.active;sheet.title='PV погодинно'
    sheet.append(['Погодинна PV-енергія інверторів']);sheet.merge_cells('A1:K1')
    sheet.append(['Джерело: архів Deye у PostgreSQL. Час UTC. Розрахована енергія окремого інвертора; часткові години не є повним фактом.'])
    sheet.merge_cells('A2:K2');sheet['A2'].alignment=Alignment(wrap_text=True);sheet.row_dimensions[2].height=34
    sheet.append([])
    sheet.append(['Дата UTC','Початок години UTC','PV-енергія, кВт·год','Покриття, с','Статус','Метод','Обсяг даних','Capture ID','SHA256','Отримано UTC','Версія мапінгу'])
    for capture in captures:
        if capture['scope']!='device_only':raise ValueError('device scope required')
        receipt=datetime.fromisoformat(capture['retrievedAtUtc']).astimezone(timezone.utc).replace(tzinfo=None)
        for hour in capture['hours']:
            start=datetime.fromisoformat(hour['hourUtc']).astimezone(timezone.utc).replace(tzinfo=None)
            status='Повна' if hour['complete'] else 'Часткова' if hour['coveredSeconds'] else 'Немає даних'
            sheet.append([start.date(),start,hour['derivedEnergyKwh'],hour['coveredSeconds'],status,
                hour['method'],'device_only',capture['captureId'],capture['documentSha256'],receipt,capture['mappingVersion']])
            line=sheet.max_row
            for cell in sheet[line]:
                if isinstance(cell.value,str):cell.data_type='s'
            sheet.cell(line,1).number_format='yyyy-mm-dd'
            for col in (2,10):sheet.cell(line,col).number_format='yyyy-mm-dd hh:mm'
            sheet.cell(line,3).number_format='0.000'
            sheet.cell(line,4).number_format='0'
            if not hour['complete']:
                sheet.cell(line,5).fill=PatternFill('solid',fgColor='FFF0C2')
    for cell in sheet[4]:
        cell.fill=PatternFill('solid',fgColor='12323D');cell.font=Font(color='FFFFFF',bold=True)
        cell.alignment=Alignment(wrap_text=True,vertical='center')
    sheet.row_dimensions[4].height=32
    sheet['A1'].font=Font(size=16,bold=True,color='12323D')
    for col,width in zip('ABCDEFGHIJK',(14,23,23,17,18,28,18,39,68,23,28)):
        sheet.column_dimensions[col].width=width
    sheet.freeze_panes='C5';sheet.auto_filter.ref=f'A4:K{max(4,sheet.max_row)}'
    sheet.sheet_view.showGridLines=False
    output=BytesIO();book.save(output);return output.getvalue()
