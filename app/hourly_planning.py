"""Pure hourly planning and export helpers for HIOS Forecast."""
from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from io import BytesIO, StringIO
import csv
import math
from typing import Iterable


@dataclass(frozen=True)
class HourlyForecast:
    interval_start_utc: datetime
    interval_end_utc: datetime
    predicted_power_kw: float
    predicted_energy_kwh: float
    quality_flags: tuple[str, ...]


@dataclass(frozen=True)
class HourlyPlan:
    interval_start_utc: datetime
    interval_end_utc: datetime
    predicted_power_kw: float
    predicted_generation_kwh: float
    consumption_kwh: float | None
    rdn_price_uah_per_kwh: float | None
    net_grid_kwh: float | None
    estimated_import_cost_uah: float | None
    estimated_export_value_uah: float | None
    quality_flags: tuple[str, ...]


def _number(value: object, field: str) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value < 0:
        raise ValueError(f"{field} must be a finite non-negative number")
    return float(value)


def hourly_plan(points: Iterable[HourlyForecast], consumption_kwh: list[float] | None = None,
                rdn_price_uah_per_kwh: list[float] | None = None) -> tuple[HourlyPlan, ...]:
    """Combine an immutable forecast with customer-supplied consumption and RDN scenarios.

    The function is intentionally deterministic and does not bid, trade, or control equipment.
    Lists must align one-for-one with forecast intervals when supplied.
    """
    values = tuple(points)
    if consumption_kwh is not None and len(consumption_kwh) != len(values):
        raise ValueError("consumptionKwh must contain one value per forecast interval")
    if rdn_price_uah_per_kwh is not None and len(rdn_price_uah_per_kwh) != len(values):
        raise ValueError("rdnPriceUahPerKwh must contain one value per forecast interval")
    result = []
    for index, point in enumerate(values):
        consumption = None if consumption_kwh is None else _number(consumption_kwh[index], "consumptionKwh")
        price = None if rdn_price_uah_per_kwh is None else _number(rdn_price_uah_per_kwh[index], "rdnPriceUahPerKwh")
        net = None if consumption is None else consumption - point.predicted_energy_kwh
        import_cost = None if net is None or price is None else max(net, 0) * price
        export_value = None if net is None or price is None else max(-net, 0) * price
        result.append(HourlyPlan(
            point.interval_start_utc, point.interval_end_utc, point.predicted_power_kw,
            point.predicted_energy_kwh, consumption, price, net, import_cost, export_value,
            point.quality_flags,
        ))
    return tuple(result)


def plan_rows(plan: Iterable[HourlyPlan]) -> list[dict[str, object]]:
    return [{
        "intervalStartUtc": item.interval_start_utc.isoformat(),
        "intervalEndUtc": item.interval_end_utc.isoformat(),
        "predictedPowerKw": item.predicted_power_kw,
        "predictedGenerationKwh": item.predicted_generation_kwh,
        "consumptionKwh": item.consumption_kwh,
        "rdnPriceUahPerKwh": item.rdn_price_uah_per_kwh,
        "netGridKwh": item.net_grid_kwh,
        "estimatedImportCostUah": item.estimated_import_cost_uah,
        "estimatedExportValueUah": item.estimated_export_value_uah,
        "qualityFlags": list(item.quality_flags),
    } for item in plan]


def csv_export(plan: Iterable[HourlyPlan]) -> str:
    output = StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(["interval_start_utc", "interval_end_utc", "predicted_power_kw", "predicted_generation_kwh",
                     "consumption_kwh", "rdn_price_uah_per_kwh", "net_grid_kwh", "estimated_import_cost_uah",
                     "estimated_export_value_uah", "quality_flags"])
    for item in plan:
        writer.writerow([item.interval_start_utc.isoformat(), item.interval_end_utc.isoformat(), item.predicted_power_kw,
                         item.predicted_generation_kwh, item.consumption_kwh, item.rdn_price_uah_per_kwh,
                         item.net_grid_kwh, item.estimated_import_cost_uah, item.estimated_export_value_uah,
                         ";".join(item.quality_flags)])
    return output.getvalue()


def xlsx_export(plan: Iterable[HourlyPlan]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    book = Workbook()
    sheet = book.active
    sheet.title = "Hourly plan"
    headers = ["Interval start UTC", "Interval end UTC", "Forecast power, kW", "Forecast generation, kWh",
               "Consumption, kWh", "RDN price, UAH/kWh", "Net grid, kWh", "Import cost, UAH",
               "Export value, UAH", "Quality flags"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173F4A")
    for item in plan:
        sheet.append([item.interval_start_utc.replace(tzinfo=None), item.interval_end_utc.replace(tzinfo=None),
                      item.predicted_power_kw, item.predicted_generation_kwh, item.consumption_kwh,
                      item.rdn_price_uah_per_kwh, item.net_grid_kwh, item.estimated_import_cost_uah,
                      item.estimated_export_value_uah, ";".join(item.quality_flags)])
    for column, width in {"A": 22, "B": 22, "C": 19, "D": 23, "E": 19, "F": 22, "G": 18, "H": 19, "I": 20, "J": 28}.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    for row in sheet.iter_rows(min_row=2, max_col=9):
        for cell in row:
            if cell.column in (1, 2):
                cell.number_format = "yyyy-mm-dd hh:mm"
            elif cell.column >= 3:
                cell.number_format = "0.00"
    data = BytesIO()
    book.save(data)
    return data.getvalue()
