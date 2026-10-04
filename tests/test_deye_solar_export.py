from io import BytesIO
from openpyxl import load_workbook
from app.deye_solar_export import solar_archive_xlsx


def test_export_preserves_blank_zero_partial_and_lineage_without_formulas():
    capture={'scope':'device_only','retrievedAtUtc':'2020-01-02T01:00:00+00:00',
        'captureId':'public-capture','documentSha256':'a'*64,'mappingVersion':'=HYPERLINK("bad")',
        'hours':[{'hourUtc':f'2020-01-01T0{h}:00:00+00:00','derivedEnergyKwh':energy,
                  'coveredSeconds':coverage,'complete':coverage==3600,'method':'linear_power_integration'}
                 for h,energy,coverage in [(0,None,0),(1,0,3600),(2,1.25,3000)]]}
    sheet=load_workbook(BytesIO(solar_archive_xlsx([capture]))).active
    assert sheet['C5'].value is None and sheet['E5'].value=='Немає даних'
    assert sheet['C6'].value==0 and sheet['E6'].value=='Повна'
    assert sheet['C7'].value==1.25 and sheet['D7'].value==3000 and sheet['E7'].value=='Часткова'
    assert sheet['H5'].value=='public-capture' and sheet['I5'].value=='a'*64
    assert sheet['K5'].data_type=='s'
    assert sheet['B5'].value.hour==0 and sheet['J5'].value.hour==1
    assert not any(cell.data_type=='f' for row in sheet for cell in row)
    assert sheet.freeze_panes=='C5'


def test_empty_archive_has_headers_without_fabricated_rows():
    sheet=load_workbook(BytesIO(solar_archive_xlsx([]))).active
    assert sheet.max_row==4 and sheet['C4'].value=='PV-енергія, кВт·год'
